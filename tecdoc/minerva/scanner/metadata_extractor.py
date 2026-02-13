"""Estrazione metadati da file docx, pdf, xlsx, pptx, eml, txt, xml, sql e altri."""

from __future__ import annotations

import email
import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


def extract_metadata(file_path: Path) -> dict[str, Any]:
    """Estrae i metadati da un file in base all'estensione.

    Restituisce un dizionario con i metadati estratti.
    I campi non disponibili sono impostati a None.
    """
    ext = file_path.suffix.lower()
    result: dict[str, Any] = {
        "autore": None,
        "autore_ultima_modifica": None,
        "versione_interna": None,
        "documento_titolo": None,
        "documento_parole_chiave": None,
        "titoli_estratti": None,
        "indice_contenuti": None,
        "sommario_estratto": None,
        "testo_contenuto": None,  # Testo per analisi AI
    }

    try:
        if ext == ".docx":
            result = _extract_docx(file_path, result)
        elif ext == ".pdf":
            result = _extract_pdf(file_path, result)
        elif ext in (".xlsx", ".xls"):
            result = _extract_excel(file_path, result)
        elif ext in (".txt", ".xml", ".sql", ".cmd", ".ksh"):
            result = _extract_txt(file_path, result)
        elif ext == ".pptx":
            result = _extract_pptx(file_path, result)
        elif ext == ".eml":
            result = _extract_eml(file_path, result)
        elif ext in (".doc", ".ppt", ".mpp"):
            logger.debug("Formato %s legacy/binario: metadati limitati per %s", ext, file_path.name)
    except Exception as e:
        logger.warning("Errore estrazione metadati da %s: %s", file_path.name, e)

    return result


def _extract_docx(file_path: Path, result: dict[str, Any]) -> dict[str, Any]:
    """Estrae metadati da file .docx usando python-docx."""
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx non installato, skip estrazione .docx")
        return result

    try:
        doc = Document(str(file_path))
        props = doc.core_properties

        result["autore"] = props.author or None
        result["autore_ultima_modifica"] = props.last_modified_by or None
        result["versione_interna"] = props.version or None
        result["documento_titolo"] = props.title or None

        if props.keywords:
            result["documento_parole_chiave"] = props.keywords

        # Estrai titoli e indice dei contenuti gerarchico
        titoli = []
        toc_lines = []
        testo_parti = []
        # Contatori per numerazione gerarchica (livello -> contatore)
        counters = [0] * 10  # Supporta fino a Heading 9

        for para in doc.paragraphs:
            style_name = para.style.name if para.style and para.style.name else ""
            text = para.text.strip()

            if style_name.startswith("Heading") and text:
                titoli.append(text)
                # Estrai livello heading (Heading 1 -> 1, Heading 2 -> 2, etc.)
                try:
                    level = int(style_name.replace("Heading", "").strip())
                except ValueError:
                    level = 1
                level = max(1, min(level, 9))

                # Aggiorna contatori: incrementa il livello corrente, azzera i successivi
                counters[level - 1] += 1
                for i in range(level, len(counters)):
                    counters[i] = 0

                # Costruisci numero gerarchico (es. "2.1.3")
                num_parts = [str(counters[i]) for i in range(level)]
                num = ".".join(num_parts)

                indent = "  " * (level - 1)
                toc_lines.append(f"{indent}{num} {text}")

            if text:
                testo_parti.append(text)

        if titoli:
            result["titoli_estratti"] = "\n".join(titoli)
        if toc_lines:
            result["indice_contenuti"] = "\n".join(toc_lines)

        result["testo_contenuto"] = "\n".join(testo_parti)

    except Exception as e:
        logger.warning("Errore lettura docx %s: %s", file_path.name, e)

    return result


def _extract_pdf(file_path: Path, result: dict[str, Any]) -> dict[str, Any]:
    """Estrae metadati da file .pdf usando PyPDF2."""
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        logger.warning("PyPDF2 non installato, skip estrazione .pdf")
        return result

    try:
        reader = PdfReader(str(file_path))
        meta = reader.metadata

        if meta:
            result["autore"] = meta.author or None
            result["documento_titolo"] = meta.title or None
            if meta.subject:
                result["documento_parole_chiave"] = meta.subject

        # Estrai indice dei contenuti dai bookmarks/outline
        try:
            outline = reader.outline
            if outline:
                toc_lines: list[str] = []
                _walk_pdf_outline(outline, toc_lines, depth=0)
                if toc_lines:
                    result["indice_contenuti"] = "\n".join(toc_lines)
        except Exception:
            pass  # Molti PDF non hanno bookmarks

        # Estrai testo dalle prime pagine
        testo_parti = []
        max_pages = min(20, len(reader.pages))
        for i in range(max_pages):
            try:
                text = reader.pages[i].extract_text()
                if text:
                    testo_parti.append(text.strip())
            except Exception:
                pass

        result["testo_contenuto"] = "\n".join(testo_parti)

    except Exception as e:
        logger.warning("Errore lettura pdf %s: %s", file_path.name, e)

    return result


def _walk_pdf_outline(outline: list, toc_lines: list[str], depth: int) -> None:
    """Percorre ricorsivamente l'outline PDF e costruisce l'indice dei contenuti."""
    for item in outline:
        if isinstance(item, list):
            # Sotto-livello: lista annidata = figli del bookmark precedente
            _walk_pdf_outline(item, toc_lines, depth + 1)
        else:
            # Bookmark: ha attributo .title
            title = getattr(item, "title", None) or str(item)
            title = title.strip()
            if title:
                indent = "  " * depth
                toc_lines.append(f"{indent}{title}")


def _extract_txt(file_path: Path, result: dict[str, Any]) -> dict[str, Any]:
    """Estrae testo da file .txt."""
    encodings = ["utf-8", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            text = file_path.read_text(encoding=enc)
            result["testo_contenuto"] = text
            # Usa la prima riga non vuota come titolo
            for line in text.splitlines():
                stripped = line.strip()
                if stripped:
                    result["documento_titolo"] = stripped[:200]
                    break
            return result
        except (UnicodeDecodeError, ValueError):
            continue
    logger.warning("Impossibile decodificare %s con encoding noti", file_path.name)
    return result


def _extract_pptx(file_path: Path, result: dict[str, Any]) -> dict[str, Any]:
    """Estrae metadati da file .pptx usando python-pptx."""
    try:
        from pptx import Presentation
    except ImportError:
        logger.warning("python-pptx non installato, skip estrazione .pptx")
        return result

    try:
        prs = Presentation(str(file_path))
        props = prs.core_properties

        result["autore"] = props.author or None
        result["autore_ultima_modifica"] = props.last_modified_by or None
        result["documento_titolo"] = props.title or None
        if props.keywords:
            result["documento_parole_chiave"] = props.keywords

        titoli = []
        testo_parti = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if not text:
                            continue
                        testo_parti.append(text)
                # Titoli dalle slide (primo placeholder di tipo titolo)
                if shape.is_placeholder and shape.placeholder_format.idx == 0:
                    title_text = shape.text.strip()
                    if title_text:
                        titoli.append(title_text)

        if titoli:
            result["titoli_estratti"] = "\n".join(titoli)
            # Indice contenuti: slide titles numerate
            toc_lines = [f"{i}. {t}" for i, t in enumerate(titoli, 1)]
            result["indice_contenuti"] = "\n".join(toc_lines)
        result["testo_contenuto"] = "\n".join(testo_parti)

    except Exception as e:
        logger.warning("Errore lettura pptx %s: %s", file_path.name, e)

    return result


def _extract_eml(file_path: Path, result: dict[str, Any]) -> dict[str, Any]:
    """Estrae metadati da file .eml (email)."""
    encodings = ["utf-8", "latin-1", "cp1252"]
    raw_bytes = file_path.read_bytes()

    for enc in encodings:
        try:
            raw_text = raw_bytes.decode(enc)
            break
        except (UnicodeDecodeError, ValueError):
            continue
    else:
        logger.warning("Impossibile decodificare %s", file_path.name)
        return result

    msg = email.message_from_string(raw_text)

    result["documento_titolo"] = msg.get("Subject") or None
    result["autore"] = msg.get("From") or None

    # Estrai corpo testuale
    testo_parti = []
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain":
                payload = part.get_payload(decode=True)
                if payload:
                    testo_parti.append(payload.decode("utf-8", errors="replace"))
    else:
        payload = msg.get_payload(decode=True)
        if payload:
            testo_parti.append(payload.decode("utf-8", errors="replace"))

    result["testo_contenuto"] = "\n".join(testo_parti)
    return result


def _extract_excel(file_path: Path, result: dict[str, Any]) -> dict[str, Any]:
    """Estrae metadati da file .xlsx/.xls."""
    ext = file_path.suffix.lower()

    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            logger.warning("openpyxl non installato, skip estrazione .xlsx")
            return result

        try:
            wb = load_workbook(str(file_path), read_only=True, data_only=True)
            props = wb.properties

            if props:
                result["autore"] = props.creator or None
                result["autore_ultima_modifica"] = props.lastModifiedBy or None
                result["documento_titolo"] = props.title or None
                if props.keywords:
                    result["documento_parole_chiave"] = props.keywords

            # Nomi dei fogli come info aggiuntiva
            sheet_names = wb.sheetnames
            if sheet_names:
                result["titoli_estratti"] = "Fogli: " + ", ".join(sheet_names)
                toc_lines = [f"{i}. {s}" for i, s in enumerate(sheet_names, 1)]
                result["indice_contenuti"] = "\n".join(toc_lines)

            wb.close()

        except Exception as e:
            logger.warning("Errore lettura xlsx %s: %s", file_path.name, e)

    elif ext == ".xls":
        try:
            import xlrd
        except ImportError:
            logger.warning("xlrd non installato, skip estrazione .xls")
            return result

        try:
            wb = xlrd.open_workbook(str(file_path))
            sheet_names = wb.sheet_names()
            if sheet_names:
                result["titoli_estratti"] = "Fogli: " + ", ".join(sheet_names)
                toc_lines = [f"{i}. {s}" for i, s in enumerate(sheet_names, 1)]
                result["indice_contenuti"] = "\n".join(toc_lines)

        except Exception as e:
            logger.warning("Errore lettura xls %s: %s", file_path.name, e)

    return result
